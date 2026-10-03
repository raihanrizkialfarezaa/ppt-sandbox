"""Baca gaya file .pptx rujukan: ringkasan font, warna, dan struktur slide.

Dipakai oleh /generate agar deck baru bisa meniru gaya referensi.
File rujukan tidak diubah sama sekali.

    python tools\\ref.py --file references\\contoh.pptx
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _rgb(run) -> str | None:
    try:
        c = run.font.color
        if c is not None and c.rgb is not None:
            return f"#{c.rgb}"
    except Exception:
        pass
    return None


def inspect_pptx(path: str) -> dict:
    from pptx import Presentation

    prs = Presentation(path)
    fonts: Counter = Counter()
    sizes: Counter = Counter()
    colors: Counter = Counter()
    fills: Counter = Counter()
    outline: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for sh in slide.shapes:
            try:
                if sh.fill is not None and sh.fill.fore_color.rgb is not None:
                    fills[f"#{sh.fill.fore_color.rgb}"] += 1
            except Exception:
                pass
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not (r.text or "").strip():
                        continue
                    if r.font.name:
                        fonts[r.font.name] += 1
                    if r.font.size:
                        sizes[int(r.font.size.pt)] += 1
                    rgb = _rgb(r)
                    if rgb:
                        colors[rgb] += 1
                    texts.append(r.text.strip())
        head = " | ".join(texts)[:100]
        outline.append(f"slide {i}: {head}")

    return {
        "file": Path(path).name,
        "slides": len(prs.slides),
        "size_in": (round(prs.slide_width.inches, 2), round(prs.slide_height.inches, 2)),
        "fonts_top": fonts.most_common(5),
        "sizes_top": sizes.most_common(6),
        "text_colors_top": colors.most_common(6),
        "fill_colors_top": fills.most_common(6),
        "outline": outline,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="Inspeksi gaya .pptx rujukan")
    ap.add_argument("--file", required=True, help="Path file .pptx rujukan")
    args = ap.parse_args()
    r = inspect_pptx(args.file)
    print(f"File: {r['file']} — {r['slides']} slide, ukuran {r['size_in'][0]}x{r['size_in'][1]} inci")
    print(f"Font dominan: {', '.join(f'{n} ({c}x)' for n, c in r['fonts_top']) or '-'}")
    print(f"Ukuran font umum: {', '.join(f'{n}pt ({c}x)' for n, c in r['sizes_top']) or '-'}")
    print(f"Warna teks umum: {', '.join(f'{n} ({c}x)' for n, c in r['text_colors_top']) or '-'}")
    print(f"Warna fill umum: {', '.join(f'{n} ({c}x)' for n, c in r['fill_colors_top']) or '-'}")
    print("Struktur:")
    for line in r["outline"]:
        print(f"  {line}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
