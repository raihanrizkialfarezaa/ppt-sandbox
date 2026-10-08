"""Builder — Brief dict/YAML -> Presentation 16:9 full-editable."""
from __future__ import annotations

from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Inches

from .assets import ensure_image
from .charts_mpl import render_mpl_chart
from .layouts import CONTENT_TYPES, LAYOUTS, add_footer, is_light_theme, l_bullets
from .theme_engine import load_theme
from .validate import validate_brief


def _prepare_slide_data(s: dict, workdir: Path, theme: dict, idx: int) -> dict:
    d = dict(s)
    # siapkan gambar: pastikan file ada (atau placeholder)
    if d.get("image"):
        src = str(d["image"])
        if not Path(src).is_absolute():
            src = str(workdir / src)
        dest = str(workdir / "output" / ".cache" / f"img-{idx}.png")
        try:
            d["image"] = ensure_image(src if Path(src).exists() else None, dest,
                                      label=str(d.get("title", "IMAGE"))[:20])
            # logo cover berlatar hitam -> jadikan transparan agar rapi di panel
            if d.get("type") == "cover" and d["image"]:
                from .assets import key_out_black
                keyed = str(workdir / "output" / ".cache" / f"logo-{idx}.png")
                d["image"] = key_out_black(d["image"], keyed)
        except Exception:
            d["image"] = None
    # siapkan fallback matplotlib bila diminta: chart_backend: mpl
    chart = d.get("chart")
    if isinstance(chart, dict) and chart.get("backend") == "mpl":
        colors = [theme["colors"]["accent"], theme["colors"]["accent2"],
                  theme["colors"]["success"], "#7C6FF0"]
        out_png = str(workdir / "output" / ".cache" / f"chart-{idx}.png")
        d["chart_image"] = render_mpl_chart(chart.get("kind", "bar"), chart.get("categories", []),
                                           chart.get("series", []), colors, out_png,
                                           title="")
    return d


def build_from_dict(data: dict, out_pptx: str, workdir: str = ".") -> dict:
    theme = load_theme(data.get("theme", "corporate-blue"), data.get("theme_override"))
    warnings = validate_brief(data)

    prs = Presentation()
    prs.slide_width = Inches(theme.get("slide_w_in", 13.33))
    prs.slide_height = Inches(theme.get("slide_h_in", 7.5))
    blank = prs.slide_layouts[6]

    slides = data.get("slides", [])
    total = len(slides)
    wd = Path(workdir)

    for i, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        t = s.get("type", "bullets")
        fn = LAYOUTS.get(t, l_bullets)
        d = _prepare_slide_data(s, wd, theme, i)
        if s.get("image"):
            _src = str(s["image"])
            if not Path(_src).is_absolute():
                _src = str(wd / _src)
            if not Path(_src).exists():
                warnings.append(f"slide {i}: gambar '{s['image']}' tidak ditemukan, dipakai placeholder.")
        try:
            fn(slide, theme, d)
        except Exception as e:  # jangan gagal total gara-gara 1 slide
            warnings.append(f"slide {i} ({s.get('type')}): {e}; fallback ke bullets.")
            l_bullets(slide, theme, {"title": s.get("title", f"Slide {i}"),
                                     "items": [{"title": str(e)}]})
        # notes editable
        if s.get("notes"):
            try:
                slide.notes_slide.placeholders[1].text = s["notes"]
            except Exception:
                pass
        base_light = is_light_theme(theme)
        cl = theme.get("content_light", base_light)
        on_dark = not (cl if t in CONTENT_TYPES else base_light)
        add_footer(slide, i, total, theme, on_dark=on_dark)

    Path(out_pptx).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_pptx)
    return {"pptx": out_pptx, "slides": total, "theme": theme["name"], "warnings": warnings}


def build_from_brief(brief_path: str, out_pptx: str, workdir: str = ".") -> dict:
    with open(brief_path, encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return build_from_dict(data, out_pptx, workdir=workdir)
